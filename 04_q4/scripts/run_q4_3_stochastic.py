"""Formal Q4-3: rolling extension of the Q4-2 joint scenario-tree LP.

At 00/06/12/18 the remaining net-load and price paths are rebuilt from paired
historical residual days. Grid purchase is common to all scenarios; battery
actions are scenario-tree recourse. Only the next block is executed on actual
net load before the next optimisation. All settlement uses realised prices.
"""
from __future__ import annotations

import copy
import json
import shutil
import sys
import time
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from scipy.optimize import linprog
from scipy.sparse import lil_matrix

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "02_q2/src"))
sys.path.insert(0, str(ROOT / "04_q4/scripts"))

import q4_3_data as data  # noqa: E402
from q2_near_optimal import hierarchical_tree_groups  # noqa: E402
from run_q4_2_comparison import generate_price_forecasts, read_prices  # noqa: E402

OUT = ROOT / "04_q4/outputs/q4_3_stochastic"
OUT.mkdir(parents=True, exist_ok=True)
NS = 12
STAGE_SLOTS = 24
TERMINAL_FLOOR = 3000.0
PLANNING_PENALTY = 4.5
EPS = 1e-4
ACTUAL_PRICE = read_prices()
PRICE_FORECASTS, _ = generate_price_forecasts(ACTUAL_PRICE, data.DATES)
PRICE_MODEL = "weekday_mean_4w"


def chosen_history(day: int, count: int = NS) -> np.ndarray:
    candidates = np.arange(max(data.MIN_FORECAST_DAY, day - 56), day)
    matching = candidates[(day - candidates) % 7 == 0][::-1]
    remaining = candidates[~np.isin(candidates, matching)][::-1]
    chosen = np.concatenate([matching, remaining])[:count]
    if len(chosen) != count:
        raise RuntimeError(f"Insufficient residual days before {data.DATES[day]}")
    return chosen


def point_price(day: int, issue: int) -> np.ndarray:
    """Q4-2 forecast with a causal intraday level correction."""
    start = data.ISSUE_STARTS[issue]
    pred = PRICE_FORECASTS[PRICE_MODEL][day, start:].copy()
    if issue == 0:
        return pred
    hist_start = max(0, start - 36)
    error = np.median(
        ACTUAL_PRICE[day, hist_start:start]
        - PRICE_FORECASTS[PRICE_MODEL][day, hist_start:start]
    )
    lead_hours = np.arange(len(pred)) * data.INTERVAL_HOURS
    return np.maximum(pred + error * np.exp(-lead_hours / 6.0), 0.0)


def paired_scenarios(day: int, issue: int) -> tuple[np.ndarray, np.ndarray]:
    """Causal paired residual paths for the remaining daily horizon."""
    start = data.ISSUE_STARTS[issue]
    chosen = chosen_history(day)
    current_net = data.FORECASTS_BY_ISSUE[issue]["net_kwh"][day]
    current_price = point_price(day, issue)
    net_paths, price_paths = [], []
    for hist_day in chosen:
        hist_net_fc = data.FORECASTS_BY_ISSUE[issue]["net_kwh"][hist_day]
        hist_net_actual = data.NET_KWH[hist_day, start:]
        hist_price_fc = point_price(hist_day, issue)
        hist_price_actual = ACTUAL_PRICE[hist_day, start:]
        net_paths.append(current_net + hist_net_actual - hist_net_fc)
        price_paths.append(np.maximum(current_price + hist_price_actual - hist_price_fc, 0.0))
    return np.asarray(net_paths), np.asarray(price_paths)


def solve_rolling_lp(
    net_paths: np.ndarray,
    price_paths: np.ndarray,
    opening_soc: float,
    baseline: np.ndarray | None,
) -> np.ndarray:
    """Q4-2 scenario-tree LP, augmented with Q3 adjustment variables."""
    ns, nt = net_paths.shape
    if price_paths.shape != (ns, nt):
        raise ValueError("Net-load and price scenarios must be paired")
    q0 = 0
    cursor = nt
    if baseline is not None:
        up0, down0 = cursor, cursor + nt
        cursor += 2 * nt
    c0, d0, s0, e0, w0 = cursor, cursor + ns*nt, cursor + 2*ns*nt, cursor + 3*ns*nt, cursor + 4*ns*nt
    nv = cursor + 5 * ns * nt
    obj = np.zeros(nv)
    mean_price = price_paths.mean(axis=0)
    if baseline is None:
        obj[q0:q0+nt] = mean_price
    else:
        obj[up0:up0+nt] = 1.5 * mean_price
        obj[down0:down0+nt] = -0.5 * mean_price
    for s in range(ns):
        obj[c0+s*nt:c0+(s+1)*nt] = EPS / ns
        obj[d0+s*nt:d0+(s+1)*nt] = EPS / ns
        obj[e0+s*nt:e0+(s+1)*nt] = PLANNING_PENALTY * price_paths[s] / ns

    groups = hierarchical_tree_groups(net_paths, STAGE_SLOTS)
    pairs = []
    for t in range(nt):
        for group in np.unique(groups[:, t]):
            members = np.flatnonzero(groups[:, t] == group)
            pairs.extend((t, int(members[0]), int(s)) for s in members[1:])
    base_rows = 2 * ns * nt + (nt if baseline is not None else 0)
    aeq = lil_matrix((base_rows + 2*len(pairs), nv))
    beq = np.zeros(base_rows + 2*len(pairs))
    row_offset = 0
    if baseline is not None:
        for t in range(nt):
            aeq[t, q0+t] = 1
            aeq[t, up0+t] = -1
            aeq[t, down0+t] = 1
            beq[t] = baseline[t]
        row_offset = nt
    for s in range(ns):
        for t in range(nt):
            row = row_offset + s*nt+t
            aeq[row, q0+t] = 1
            aeq[row, c0+s*nt+t] = -1
            aeq[row, d0+s*nt+t] = 1
            aeq[row, e0+s*nt+t] = 1
            aeq[row, w0+s*nt+t] = -1
            beq[row] = net_paths[s, t]
            row = row_offset + ns*nt + s*nt+t
            aeq[row, s0+s*nt+t] = 1
            aeq[row, c0+s*nt+t] = -0.9
            aeq[row, d0+s*nt+t] = 1/0.9
            if t:
                aeq[row, s0+s*nt+t-1] = -1
            else:
                beq[row] = opening_soc
    for i, (t, ref, s) in enumerate(pairs):
        row = base_rows + 2*i
        aeq[row, c0+s*nt+t] = 1
        aeq[row, c0+ref*nt+t] = -1
        aeq[row+1, d0+s*nt+t] = 1
        aeq[row+1, d0+ref*nt+t] = -1
    bounds = [(0, None)] * nt
    if baseline is not None:
        bounds += [(0, None)] * (2*nt)
    bounds += (
        [(0, data.BATTERY_INTERVAL_LIMIT_KWH)] * (2*ns*nt)
        + [(data.SOC_MIN_KWH, data.SOC_MAX_KWH)] * (ns*nt)
        + [(0, None)] * (2*ns*nt)
    )
    aub = lil_matrix((ns, nv))
    bub = np.full(ns, -TERMINAL_FLOOR)
    for s in range(ns):
        aub[s, s0+(s+1)*nt-1] = -1
    solved = linprog(
        obj, A_ub=aub.tocsr(), b_ub=bub,
        A_eq=aeq.tocsr(), b_eq=beq, bounds=bounds, method="highs",
    )
    if not solved.success:
        raise RuntimeError(solved.message)
    return solved.x[q0:q0+nt]


def execute_block(purchase: np.ndarray, actual_net: np.ndarray, opening_soc: float) -> dict:
    n = len(purchase)
    out = {name: np.zeros(n) for name in ("charge", "discharge", "emergency", "spill", "soc")}
    soc = float(opening_soc)
    for t in range(n):
        margin = float(purchase[t] - actual_net[t])
        if margin >= 0:
            out["charge"][t] = min(data.BATTERY_INTERVAL_LIMIT_KWH, margin, (data.SOC_MAX_KWH-soc)/0.9)
            out["spill"][t] = margin - out["charge"][t]
            soc += 0.9 * out["charge"][t]
        else:
            need = -margin
            out["discharge"][t] = min(data.BATTERY_INTERVAL_LIMIT_KWH, need, (soc-data.SOC_MIN_KWH)*0.9)
            out["emergency"][t] = need - out["discharge"][t]
            soc -= out["discharge"][t] / 0.9
        out["soc"][t] = soc
    out["soc_close"] = soc
    return out


def simulate(update_hours: tuple[int, ...] = (6, 12, 18)) -> dict:
    nd, nt = len(data.DATES), data.INTERVALS_PER_DAY
    names = ("plan", "adjusted", "charge", "discharge", "emergency", "spill", "soc")
    full = {name: np.zeros((nd, nt)) for name in names}
    soc_open = np.zeros(nd)
    current_soc = data.SOC_TARGET_KWH
    # January uses a causal previous-day net-load purchase to propagate SOC.
    for day in range(31):
        purchase = np.maximum(data.NET_KWH[max(day-1, 0)], 0)
        rec = execute_block(purchase, data.NET_KWH[day], current_soc)
        current_soc = float(rec["soc_close"])
    for count, day in enumerate(data.EVAL_DAYS, 1):
        soc_open[day] = current_soc
        net0, price0 = paired_scenarios(day, 0)
        day_plan = solve_rolling_lp(net0, price0, current_soc, None)
        full["plan"][day] = day_plan
        for block, issue in enumerate(data.ISSUE_HOURS):
            start, end = data.BLOCK_BOUNDS[block], data.BLOCK_BOUNDS[block+1]
            if issue == 0:
                remaining = day_plan
            elif issue in update_hours:
                net_paths, price_paths = paired_scenarios(day, issue)
                remaining = solve_rolling_lp(
                    net_paths, price_paths, current_soc, day_plan[start:]
                )
            else:
                remaining = day_plan[start:]
            purchase = remaining[:end-start]
            rec = execute_block(purchase, data.NET_KWH[day, start:end], current_soc)
            full["adjusted"][day, start:end] = purchase
            for name in ("charge", "discharge", "emergency", "spill", "soc"):
                full[name][day, start:end] = rec[name]
            current_soc = float(rec["soc_close"])
        if count % 25 == 0:
            print(f"days {count}/334", flush=True)
    idx = data.EVAL_DAYS
    result = {name: values[idx] for name, values in full.items()}
    result["soc_open"] = soc_open[idx]
    result["soc_close"] = result["soc"][:, -1]
    up = np.maximum(result["adjusted"]-result["plan"], 0)
    down = np.maximum(result["plan"]-result["adjusted"], 0)
    prices = ACTUAL_PRICE[idx]
    result["upward"], result["downward"] = up, down
    result["plan_cost"] = np.sum(prices*result["plan"], axis=1)
    result["adjustment_cost"] = np.sum(prices*(1.5*up-0.5*down), axis=1)
    result["emergency_cost"] = np.sum(5*prices*result["emergency"], axis=1)
    result["total_cost"] = result["plan_cost"]+result["adjustment_cost"]+result["emergency_cost"]
    result["metrics"] = {
        "total_cost_yuan": float(result["total_cost"].sum()),
        "plan_cost_yuan": float(result["plan_cost"].sum()),
        "adjustment_cost_yuan": float(result["adjustment_cost"].sum()),
        "emergency_cost_yuan": float(result["emergency_cost"].sum()),
        "emergency_kwh": float(result["emergency"].sum()),
        "plan_kwh": float(result["plan"].sum()),
        "final_kwh": float(result["adjusted"].sum()),
        "min_soc_kwh": float(result["soc"].min()),
        "max_soc_kwh": float(result["soc"].max()),
    }
    return result


def block_totals(x: np.ndarray) -> np.ndarray:
    return x.reshape(6, 24).sum(axis=1)


def emergency_windows(x: np.ndarray) -> list[tuple[str, float]]:
    active = np.flatnonzero(x > 1e-7)
    if not len(active):
        return [("None", 0.0)]
    groups, start, prev = [], int(active[0]), int(active[0])
    for value in active[1:]:
        value = int(value)
        if value != prev + 1:
            groups.append((f"{data.INTERVAL_LABELS[start].split('-')[0]}-{data.INTERVAL_LABELS[prev].split('-')[1]}", float(x[start:prev+1].sum())))
            start = value
        prev = value
    groups.append((f"{data.INTERVAL_LABELS[start].split('-')[0]}-{data.INTERVAL_LABELS[prev].split('-')[1]}", float(x[start:prev+1].sum())))
    return groups


def export_workbook(result: dict) -> Path:
    dest = OUT / "result4-3.xlsx"
    shutil.copy2(data.TEMPLATE_FILE, dest)
    wb = load_workbook(dest)
    plan_ws, adj_ws, bat_ws, em_ws = [wb[name] for name in (
        data.PLAN_SHEET_NAME, data.ADJUSTMENT_SHEET_NAME,
        data.BATTERY_SHEET_NAME, data.EMERGENCY_SHEET_NAME,
    )]
    for r, _date in enumerate(data.REPORT_DAY_INDEX, 2):
        i = r-2
        for c, value in enumerate(result["plan"][i], 2): plan_ws.cell(r,c).value = round(float(value), 4)
        plan_ws.cell(r,146).value = round(float(result["plan"][i].sum()), 4)
        plan_ws.cell(r,147).value = round(float(result["plan_cost"][i]), 4)
        for c, value in enumerate(result["adjusted"][i], 2): adj_ws.cell(r,c).value = round(float(value), 4)
        adj_ws.cell(r,146).value = round(float(result["adjusted"][i].sum()), 4)
        adj_ws.cell(r,147).value = round(float(result["adjustment_cost"][i]), 4)
    bat_styles = [[copy.copy(bat_ws.cell(r,c)._style) for c in range(1,7)] for r in range(2,8)]
    bat_ws.delete_rows(2, bat_ws.max_row-1)
    for i, date in enumerate(data.REPORT_DAY_INDEX):
        ch, dch = block_totals(result["charge"][i]), block_totals(result["discharge"][i])
        for b, label in enumerate(data.FOUR_HOUR_LABELS):
            r = 2+i*6+b
            for c in range(1,7): bat_ws.cell(r,c)._style = copy.copy(bat_styles[b][c-1])
            if b == 0:
                bat_ws.cell(r,1).value = date.to_pydatetime(); bat_ws.cell(r,5).value = "0:00"; bat_ws.cell(r,6).value = round(float(result["soc_open"][i]),4)
            elif b == 1:
                bat_ws.cell(r,5).value = "24:00"; bat_ws.cell(r,6).value = round(float(result["soc_close"][i]),4)
            bat_ws.cell(r,2).value = label; bat_ws.cell(r,3).value = round(float(ch[b]),4); bat_ws.cell(r,4).value = round(float(dch[b]),4)
    em_styles = [[copy.copy(em_ws.cell(r,c)._style) for c in range(1,4)] for r in (2,3,4)]
    em_ws.delete_rows(2, em_ws.max_row-1)
    r = 2
    for i, date in enumerate(data.REPORT_DAY_INDEX):
        windows = emergency_windows(result["emergency"][i])
        for j, (window, energy) in enumerate(windows):
            style = 0 if j == 0 else (2 if j == len(windows)-1 else 1)
            for c in range(1,4): em_ws.cell(r,c)._style = copy.copy(em_styles[style][c-1])
            if j == 0: em_ws.cell(r,1).value = date.to_pydatetime()
            em_ws.cell(r,2).value = window; em_ws.cell(r,3).value = round(float(energy),4); r += 1
    wb.save(dest); wb.close()
    return dest


def main() -> None:
    tic = time.perf_counter()
    result = simulate()
    np.savez_compressed(OUT / "policy.npz", **{k:v for k,v in result.items() if isinstance(v,np.ndarray)})
    daily = pd.DataFrame({
        "date": data.REPORT_DAY_INDEX,
        "plan_cost_yuan": result["plan_cost"],
        "adjustment_cost_yuan": result["adjustment_cost"],
        "emergency_cost_yuan": result["emergency_cost"],
        "total_cost_yuan": result["total_cost"],
        "emergency_kwh": result["emergency"].sum(axis=1),
        "soc_open_kwh": result["soc_open"],
        "soc_close_kwh": result["soc_close"],
    })
    daily.to_csv(OUT / "daily_results.csv", index=False)
    workbook = export_workbook(result)
    summary = {
        "method": "Q4-2 paired net-load/price residual scenario tree with rolling re-optimisation at 0/6/12/18",
        "price_model": PRICE_MODEL,
        "scenario_count": NS,
        "terminal_floor_kwh": TERMINAL_FLOOR,
        "planning_emergency_multiplier": PLANNING_PENALTY,
        "metrics": result["metrics"],
        "workbook": str(workbook),
        "elapsed_s": time.perf_counter()-tic,
    }
    (OUT / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2), flush=True)


if __name__ == "__main__":
    main()
