"""Perfect-information lower bound for Q3 (God's-eye hindsight LP).

Q2 paper approach: if the realized net-load path is known, the planner
commits only the energy that is actually used after optimal storage
arbitrage, pays the normal tariff, and never incurs 0.5 / 1.5 / 5
mismatch or emergency charges.

Does not change v002. Local only; do not push.
"""
from __future__ import annotations

import json
import time
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from scipy.optimize import linprog
from scipy.sparse import coo_matrix

HERE = Path(__file__).resolve().parent
Q3_DIR = HERE.parent
ROOT = Q3_DIR.parent
DATA_DIR = ROOT / "附件"
OUT = Q3_DIR / "outputs" / "perfect_info"
OUT.mkdir(parents=True, exist_ok=True)

INTERVALS = 144
DT = 1.0 / 6.0
SOC_MIN = 1_200.0
SOC_MAX = 10_800.0
SOC0 = 6_000.0
PMAX = 5_000.0 * DT
ETA_C = 0.9
ETA_D = 0.9
TIE = 1e-7
V002 = 13_521_409.707452236
V001 = 14_438_642.683144223
V000 = 14_669_597.26287158
EVAL_START = 31  # 2025-02-01
REPORT_DATES = ("2025-03-20", "2025-06-21", "2025-09-23", "2025-12-21")


def load_arrays() -> dict:
    price = pd.to_numeric(pd.read_excel(DATA_DIR / "附件1.xlsx").iloc[:, 1], errors="raise").to_numpy(float)
    book = pd.ExcelFile(DATA_DIR / "附件2.xlsx")
    load_frame = book.parse(book.sheet_names[0])
    pv_frame = book.parse(book.sheet_names[1])
    dates = pd.DatetimeIndex(pd.to_datetime(load_frame.iloc[:, 0])).normalize()
    load_kwh = load_frame.iloc[:, 1:].to_numpy(float) * DT
    pv_kwh = pv_frame.iloc[:, 1:].to_numpy(float) * DT
    net = load_kwh - pv_kwh
    template = DATA_DIR / "附件5" / "result3.xlsx"
    wb = load_workbook(template, read_only=True, data_only=False)
    plan_sheet = wb[wb.sheetnames[0]]
    labels = [str(plan_sheet.cell(1, c).value).strip() for c in range(2, 146)]
    report_index = pd.DatetimeIndex(
        [pd.Timestamp(plan_sheet.cell(r, 1).value).normalize() for r in range(2, 336)]
    )
    wb.close()
    assert price.shape == (INTERVALS,)
    assert net.shape == (365, INTERVALS)
    assert np.array_equal(report_index, dates[EVAL_START:])
    return {
        "price": price,
        "dates": dates,
        "net": net,
        "labels": labels,
        "report_index": report_index,
    }


def solve_pi(net_flat: np.ndarray, price_flat: np.ndarray, opening_soc: float) -> dict:
    n = int(net_flat.size)
    ng, nc, nd, nw, nsoc = 0, n, 2 * n, 3 * n, 4 * n
    nv = 5 * n
    t = np.arange(n)
    obj = np.zeros(nv)
    obj[ng:nc] = price_flat
    obj[nc:nd] = TIE
    obj[nd:nw] = TIE
    obj[nw:nsoc] = TIE

    rows = np.concatenate([t, t, t, t, n + t, n + t, n + t, n + t[1:]])
    cols = np.concatenate(
        [ng + t, nd + t, nc + t, nw + t, nsoc + t, nc + t, nd + t, nsoc + t[1:] - 1]
    )
    data = np.concatenate(
        [
            np.ones(n),
            np.ones(n),
            -np.ones(n),
            -np.ones(n),
            np.ones(n),
            np.full(n, -ETA_C),
            np.full(n, 1.0 / ETA_D),
            -np.ones(n - 1),
        ]
    )
    a_eq = coo_matrix((data, (rows, cols)), shape=(2 * n, nv)).tocsr()
    b_eq = np.concatenate([net_flat, np.zeros(n)])
    b_eq[n] = opening_soc
    bounds = (
        [(0.0, None)] * n
        + [(0.0, PMAX)] * n
        + [(0.0, PMAX)] * n
        + [(0.0, None)] * n
        + [(SOC_MIN, SOC_MAX)] * n
    )
    t0 = time.perf_counter()
    sol = linprog(obj, A_eq=a_eq, b_eq=b_eq, bounds=bounds, method="highs")
    elapsed = time.perf_counter() - t0
    if not sol.success:
        raise RuntimeError(sol.message)
    x = sol.x
    purchase = x[ng:nc]
    charge = x[nc:nd]
    discharge = x[nd:nw]
    spill = x[nw:nsoc]
    soc = x[nsoc:]
    residual = purchase + discharge - charge - spill - net_flat
    if np.max(np.abs(residual)) > 1e-6:
        raise RuntimeError(f"energy balance residual {np.max(np.abs(residual))}")
    paid = float(np.dot(price_flat, purchase))
    return {
        "purchase": purchase,
        "charge": charge,
        "discharge": discharge,
        "spill": spill,
        "soc": soc,
        "paid": paid,
        "solver_fun": float(sol.fun),
        "elapsed_s": elapsed,
        "n_slots": n,
        "max_balance_abs": float(np.max(np.abs(residual))),
        "min_soc": float(soc.min()),
        "max_soc": float(soc.max()),
        "end_soc": float(soc[-1]),
    }


def reshape(vec: np.ndarray, n_days: int) -> np.ndarray:
    return vec.reshape(n_days, INTERVALS)


def eval_metrics(data: dict, result: dict, day0: int, n_days: int) -> dict:
    price = data["price"]
    dates = data["dates"][day0 : day0 + n_days]
    purchase = reshape(result["purchase"], n_days)
    charge = reshape(result["charge"], n_days)
    discharge = reshape(result["discharge"], n_days)
    spill = reshape(result["spill"], n_days)
    soc = reshape(result["soc"], n_days)
    local = dates >= data["report_index"][0]
    p = purchase[local]
    daily_cost = p @ price
    eval_pos = np.flatnonzero(local)
    opens = np.empty(eval_pos.size)
    for i, pos in enumerate(eval_pos):
        if pos == 0:
            opens[i] = float(soc[0, 0] - ETA_C * charge[0, 0] + discharge[0, 0] / ETA_D)
        else:
            opens[i] = float(soc[pos - 1, -1])
    daily = pd.DataFrame(
        {
            "date": dates[local].strftime("%Y-%m-%d"),
            "soc_open_kwh": opens,
            "soc_close_kwh": soc[local, -1],
            "purchase_kwh": p.sum(axis=1),
            "charge_kwh": charge[local].sum(axis=1),
            "discharge_kwh": discharge[local].sum(axis=1),
            "spill_kwh": spill[local].sum(axis=1),
            "cost_yuan": daily_cost,
        }
    )
    return {
        "total_cost": float(daily_cost.sum()),
        "purchase_kwh": float(p.sum()),
        "charge_kwh": float(charge[local].sum()),
        "discharge_kwh": float(discharge[local].sum()),
        "spill_kwh": float(spill[local].sum()),
        "mean_soc_open": float(daily["soc_open_kwh"].mean()),
        "mean_soc_close": float(daily["soc_close_kwh"].mean()),
        "feb1_soc_open": float(daily["soc_open_kwh"].iloc[0]),
        "dec31_soc_close": float(daily["soc_close_kwh"].iloc[-1]),
        "daily": daily,
        "purchase": p,
        "soc": soc[local],
        "charge": charge[local],
        "discharge": discharge[local],
    }


def no_battery_cost(net: np.ndarray, price: np.ndarray) -> float:
    take = np.maximum(net, 0.0)
    return float((take * price[None, :]).sum())


def main() -> None:
    data = load_arrays()
    price = data["price"]
    net = data["net"]
    print("Loaded", len(data["dates"]), "days. Building perfect-information LPs.")

    cases = []

    # Case A: true problem lower bound — 1 Jan SOC=6000, operate all year, score Feb–Dec.
    net_year = net.ravel()
    price_year = np.tile(price, 365)
    print("Solving 365-day hindsight LP, S0=6000 ...")
    year = solve_pi(net_year, price_year, SOC0)
    m_year = eval_metrics(data, year, 0, 365)
    print(
        f"  year LP {year['elapsed_s']:.1f}s  eval cost={m_year['total_cost']:,.2f}  "
        f"Feb1 SOC={m_year['feb1_soc_open']:.1f}  Dec31 SOC={m_year['dec31_soc_close']:.1f}"
    )
    cases.append(("jan1_6000_fullyear", year, m_year, 0, 365))

    # Case B: evaluation window only, Feb 1 starts at 6000.
    net_eval = net[EVAL_START:].ravel()
    price_eval = np.tile(price, 365 - EVAL_START)
    print("Solving 334-day hindsight LP, Feb1 S0=6000 ...")
    ev = solve_pi(net_eval, price_eval, SOC0)
    m_ev = eval_metrics(data, ev, EVAL_START, 365 - EVAL_START)
    print(
        f"  eval LP {ev['elapsed_s']:.1f}s  cost={m_ev['total_cost']:,.2f}  "
        f"Feb1 SOC={m_ev['feb1_soc_open']:.1f}  Dec31 SOC={m_ev['dec31_soc_close']:.1f}"
    )
    cases.append(("feb1_6000_evalonly", ev, m_ev, EVAL_START, 365 - EVAL_START))

    nb = no_battery_cost(net[EVAL_START:], price)
    print(f"No-battery (buy every positive net slot): {nb:,.2f}")

    official = m_year  # Q2-style: common Jan-1 init, January is warmup, score the required dates
    lb = official["total_cost"]
    summary = {
        "keeps_champion": "v002",
        "v002_yuan": V002,
        "theoretical_lower_bound_yuan": lb,
        "gap_v002_minus_lb_yuan": V002 - lb,
        "gap_v002_pct": 100.0 * (V002 - lb) / V002,
        "gap_v001_minus_lb_yuan": V001 - lb,
        "gap_v000_minus_lb_yuan": V000 - lb,
        "no_battery_yuan": nb,
        "eval_period": "2025-02-01 to 2025-12-31",
        "official_case": "jan1_6000_fullyear",
        "meaning": (
            "Know every 10-min realized load and PV from 1 Jan. Choose purchase and "
            "battery jointly. Pay only the normal tariff. No 0.5 cancel, 1.5 extra, or 5x emergency. "
            "January cost is excluded; January storage may be used to position 1 Feb. "
            "No salvage on 31 Dec. This is a lower bound on any feasible Q3 policy that starts at 6000 kWh on 1 Jan."
        ),
        "q2_paper_analog": (
            "Same construction as the Q2 perfect-information figure 12,226,852.42 CNY, "
            "which the Q2 manuscript withdrew only because the script was missing. "
            "Q3 settlement collapses to p·G under hindsight because plan=take and E=0."
        ),
        "cases": {},
        "paper_days": {},
    }
    for name, raw, met, day0, n_days in cases:
        summary["cases"][name] = {
            "total_cost_yuan": met["total_cost"],
            "purchase_kwh": met["purchase_kwh"],
            "charge_kwh": met["charge_kwh"],
            "discharge_kwh": met["discharge_kwh"],
            "spill_kwh": met["spill_kwh"],
            "feb1_soc_open": met["feb1_soc_open"],
            "dec31_soc_close": met["dec31_soc_close"],
            "mean_soc_open": met["mean_soc_open"],
            "mean_soc_close": met["mean_soc_close"],
            "solver_elapsed_s": raw["elapsed_s"],
            "max_balance_abs": raw["max_balance_abs"],
            "min_soc": raw["min_soc"],
            "max_soc": raw["max_soc"],
        }
        met["daily"].to_csv(OUT / f"daily_{name}.csv", index=False, encoding="utf-8-sig")
        monthly = met["daily"].copy()
        monthly["month"] = pd.to_datetime(monthly["date"]).dt.strftime("%Y-%m")
        monthly.groupby("month", as_index=False)["cost_yuan"].sum().to_csv(
            OUT / f"monthly_{name}.csv", index=False, encoding="utf-8-sig"
        )

    dates_eval = data["report_index"]
    for date in REPORT_DATES:
        row = int(np.flatnonzero(dates_eval == pd.Timestamp(date))[0])
        summary["paper_days"][date] = {
            "cost_yuan": float(official["daily"]["cost_yuan"].iloc[row]),
            "purchase_kwh": float(official["daily"]["purchase_kwh"].iloc[row]),
            "soc_open": float(official["daily"]["soc_open_kwh"].iloc[row]),
            "soc_close": float(official["daily"]["soc_close_kwh"].iloc[row]),
            "charge_kwh": float(official["daily"]["charge_kwh"].iloc[row]),
            "discharge_kwh": float(official["daily"]["discharge_kwh"].iloc[row]),
        }

    (OUT / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    official["daily"].to_csv(OUT / "daily_eval.csv", index=False, encoding="utf-8-sig")
    print("\n===== THEORETICAL LOWER BOUND =====")
    print(f"  {lb:,.2f} yuan")
    print(f"  v002 {V002:,.2f}  gap {V002-lb:,.2f} ({summary['gap_v002_pct']:.2f}%)")
    print(f"  v001 {V001:,.2f}")
    print(f"  v000 {V000:,.2f}")
    print(f"  no battery {nb:,.2f}")
    print("Written", OUT)


if __name__ == "__main__":
    main()
