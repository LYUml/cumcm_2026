"""Independent validation checks for q2_model.py and its exported workbook."""

from __future__ import annotations

import json
import sys
from pathlib import Path

import numpy as np
import pandas as pd

from q2_model import (
    DT_HOURS,
    SOC_INITIAL,
    export_result2,
    generate_causal_forecasts,
    read_inputs,
    settle_plan,
    solve_day_ahead_lp,
)

_BUNDLED_SITE = Path(
    "/Users/lml/.cache/codex-runtimes/codex-primary-runtime/dependencies/"
    "python/lib/python3.12/site-packages"
)
if _BUNDLED_SITE.exists():
    sys.path.append(str(_BUNDLED_SITE))
from openpyxl import load_workbook  # noqa: E402


def main() -> None:
    root = Path(__file__).resolve().parent
    out = root / "q2_outputs"
    dates, load, pv, price = read_inputs(root)
    net = load - pv

    comparison = pd.read_csv(out / "model_comparison.csv")
    best = str(comparison.iloc[0]["model"])
    assert best.startswith("scenario_")
    assert comparison["selection_cost_Feb_Aug_yuan"].is_monotonic_increasing

    # Causality test: changing all data from a cutoff onward cannot alter any
    # forecast issued before that cutoff.
    cutoff = 120
    base = generate_causal_forecasts(net[:cutoff], dates[:cutoff])
    altered = net.copy()
    altered[cutoff:] += 1_000_000.0
    changed = generate_causal_forecasts(altered, dates)
    leakage_max_difference = max(
        float(np.max(np.abs(base[name] - changed[name][: cutoff - 31])))
        for name in base
    )
    assert leakage_max_difference == 0.0

    # Independently compute the perfect-information lower bound using the same
    # physical LP. Any causal forecast policy must cost at least this much.
    perfect_cost = 0.0
    for day in range(31, len(net)):
        perfect_plan = solve_day_ahead_lp(net[day], price, SOC_INITIAL)
        perfect_settlement = settle_plan(perfect_plan, net[day], price)
        assert perfect_settlement["emergency"].sum() <= 1e-6
        perfect_cost += perfect_settlement["total_cost"]
    best_cost = float(comparison.iloc[0]["total_cost_yuan"])
    assert best_cost + 1e-6 >= perfect_cost

    # Reopen the workbook and reconcile its plan totals with the comparison CSV.
    wb = load_workbook(out / "result2.xlsx", read_only=True, data_only=True)
    plan_ws = wb["计划购电量"]
    battery_ws = wb["充放电量"]
    emergency_ws = wb["紧急购电量"]
    assert plan_ws.max_row == 335 and plan_ws.max_column == 147
    assert battery_ws.max_row == 1 + 334 * 6
    assert emergency_ws.max_row > 334
    workbook_plan_cost = sum(
        float(plan_ws.cell(row=row, column=147).value) for row in range(2, 336)
    )
    expected_plan_cost = float(comparison.iloc[0]["planned_cost_yuan"])
    assert abs(workbook_plan_cost - expected_plan_cost) < 0.01
    wb.close()

    report = {
        "status": "passed",
        "best_model": best,
        "leakage_max_difference": leakage_max_difference,
        "perfect_information_lower_bound_yuan": perfect_cost,
        "best_model_total_cost_yuan": best_cost,
        "optimality_gap_vs_perfect_information_percent": 100 * (best_cost / perfect_cost - 1),
        "workbook_plan_cost_reconciliation_error_yuan": workbook_plan_cost - expected_plan_cost,
        "workbook_plan_rows": 334,
        "workbook_battery_rows": 334 * 6,
    }
    with open(out / "independent_validation.json", "w", encoding="utf-8") as f:
        json.dump(report, f, ensure_ascii=False, indent=2)
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
