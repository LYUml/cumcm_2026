"""Causal Q4-3 data preparation, independent of the archived Q3 solution."""
from __future__ import annotations

from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
APPENDIX = ROOT / "00_problem/appendix"
TEMPLATE_FILE = APPENDIX / "附件5/result3.xlsx"

INTERVALS_PER_DAY = 144
INTERVAL_HOURS = 1/6
BATTERY_INTERVAL_LIMIT_KWH = 5000/6
SOC_MIN_KWH, SOC_MAX_KWH, SOC_TARGET_KWH = 1200.0, 10800.0, 6000.0
MIN_FORECAST_DAY = 14
ISSUE_HOURS = (0, 6, 12, 18)
REPORT_DATES = ("2025-03-20", "2025-06-21", "2025-09-23", "2025-12-21")
REPORT_INTERVALS = ("10:00-10:10", "12:00-12:10", "14:00-14:10", "16:00-16:10", "18:00-18:10", "20:00-20:10")
FOUR_HOUR_LABELS = ("0:00-4:00", "4:00-8:00", "8:00-12:00", "12:00-16:00", "16:00-20:00", "20:00-24:00")

book = load_workbook(TEMPLATE_FILE, read_only=True, data_only=False)
PLAN_SHEET_NAME, ADJUSTMENT_SHEET_NAME, BATTERY_SHEET_NAME, EMERGENCY_SHEET_NAME = book.sheetnames
ws = book[PLAN_SHEET_NAME]
INTERVAL_LABELS = [str(ws.cell(1, c).value).strip() for c in range(2, 146)]
REPORT_DAY_INDEX = pd.DatetimeIndex([pd.Timestamp(ws.cell(r, 1).value).normalize() for r in range(2, 336)])
book.close()

SLOT_END_HOURS = np.arange(1, 145) / 6
ISSUE_STARTS = {h: (0 if h == 0 else int(np.flatnonzero(np.isclose(SLOT_END_HOURS, h))[0])) for h in ISSUE_HOURS}
BLOCK_BOUNDS = [ISSUE_STARTS[h] for h in ISSUE_HOURS] + [144]

history = pd.ExcelFile(APPENDIX / "附件2.xlsx")
load_frame = history.parse(history.sheet_names[0])
pv_frame = history.parse(history.sheet_names[1])
DATES = pd.DatetimeIndex(pd.to_datetime(load_frame.iloc[:, 0])).normalize()
LOAD_KW = load_frame.iloc[:, 1:].to_numpy(float)
PV_KW = pv_frame.iloc[:, 1:].to_numpy(float)
NET_KWH = (LOAD_KW - PV_KW) * INTERVAL_HOURS
EVAL_DAYS = np.arange(31, 365)
DAY_OF_WEEK = DATES.dayofweek.to_numpy()

forecast_frame = pd.read_excel(APPENDIX / "附件3.xlsx")
forecast_dates = pd.to_datetime(forecast_frame.iloc[:, 0].replace("", np.nan).ffill()).dt.normalize()
forecast_hours = forecast_frame.iloc[:, 1].astype(str).str.extract(r"(\d+):", expand=False).astype(int)
date_to_day = {date: i for i, date in enumerate(DATES)}
PV_HOURLY = np.full((365, 4, 24), np.nan)
for r in range(len(forecast_frame)):
    PV_HOURLY[date_to_day[pd.Timestamp(forecast_dates.iloc[r])], ISSUE_HOURS.index(int(forecast_hours.iloc[r]))] = forecast_frame.iloc[r, 2:26].to_numpy(float)


def load_forecast(day: int, issue: int) -> np.ndarray:
    same = np.flatnonzero(DAY_OF_WEEK[:day] == DAY_OF_WEEK[day])[-4:]
    hist = same if len(same) else np.arange(max(0, day-7), day)
    weights = np.arange(1, len(hist)+1, dtype=float); weights /= weights.sum()
    base = np.average(LOAD_KW[hist], axis=0, weights=weights)
    start = ISSUE_STARTS[issue]
    remaining = base[start:].copy()
    if issue:
        lo = max(0, start-12)
        ratio = LOAD_KW[day, lo:start].sum() / max(base[lo:start].sum(), 1e-9)
        ratio = float(np.clip(ratio, 0.8, 1.2))
        remaining *= 1 + (ratio-1)*np.exp(-(SLOT_END_HOURS[start:]-issue)/6)
    return remaining


def pv_forecast(day: int, issue: int) -> np.ndarray:
    start = ISSUE_STARTS[issue]
    anchor = PV_KW[day-1, -1] if issue == 0 else PV_KW[day, start]
    nodes = np.concatenate([[anchor], PV_HOURLY[day, ISSUE_HOURS.index(issue)]])
    target = SLOT_END_HOURS[start:] - issue
    return np.clip(np.interp(target, np.arange(25), nodes), 0, None)


FORECASTS_BY_ISSUE = {}
for issue in ISSUE_HOURS:
    horizon = 144 - ISSUE_STARTS[issue]
    net = np.full((365, horizon), np.nan)
    for day in range(MIN_FORECAST_DAY, 365):
        net[day] = (load_forecast(day, issue)-pv_forecast(day, issue))*INTERVAL_HOURS
    FORECASTS_BY_ISSUE[issue] = {"net_kwh": net}

assert LOAD_KW.shape == PV_KW.shape == (365, 144)
assert np.isfinite(PV_HOURLY).all()
