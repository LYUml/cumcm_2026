"""Reopen the exported Q4-2 workbook and reconcile it to the model payload."""
from pathlib import Path
import json
import numpy as np
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
book_path = ROOT / "04_q4/outputs/final/result4-2.xlsx"
payload = json.loads((ROOT / "04_q4/outputs/final/result4_2_payload.json").read_text())
model = json.loads((ROOT / "04_q4/outputs/final/result4_2_model_validation.json").read_text())
book = load_workbook(book_path, read_only=False, data_only=False)
assert book.sheetnames == ["计划购电量", "充放电量", "紧急购电量"]

plan = book["计划购电量"]
values = np.array([[plan.cell(r, c).value for c in range(2, 146)] for r in range(2, 336)], float)
totals = np.array([plan.cell(r, 146).value for r in range(2, 336)], float)
costs = np.array([plan.cell(r, 147).value for r in range(2, 336)], float)
source = np.array([row[1:145] for row in payload["plan"]], float)

battery = book["充放电量"]
emergency = book["紧急购电量"]
emergency_values = np.array([emergency.cell(r, 3).value or 0 for r in range(2, 432)], float)
errors = []
tokens = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!")
for sheet in book:
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and any(token in cell.value for token in tokens):
                errors.append((sheet.title, cell.coordinate, cell.value))

checks = {
    "sheet_names": book.sheetnames,
    "plan_shape": list(values.shape),
    "max_plan_export_difference_kwh": float(np.max(np.abs(values - source))),
    "max_plan_row_total_error_kwh": float(np.max(np.abs(totals - values.sum(axis=1)))),
    "exported_plan_cost_yuan": float(costs.sum()),
    "model_plan_cost_yuan": float(model["planned_yuan"]),
    "plan_cost_difference_yuan": float(abs(costs.sum() - model["planned_yuan"])),
    "battery_populated_rows": int(sum(battery.cell(r, 3).value is not None for r in range(2, 2006))),
    "expected_battery_rows": 2004,
    "emergency_populated_rows": int(sum(emergency.cell(r, 3).value is not None for r in range(2, 432))),
    "expected_emergency_rows": 430,
    "exported_emergency_kwh": float(emergency_values.sum()),
    "model_emergency_kwh": float(model["emergency_kwh"]),
    "formula_error_count": len(errors),
}
assert checks["max_plan_export_difference_kwh"] < 1e-9
assert checks["max_plan_row_total_error_kwh"] < 1e-3
assert checks["plan_cost_difference_yuan"] < 1e-3
assert checks["battery_populated_rows"] == checks["expected_battery_rows"]
assert checks["emergency_populated_rows"] == checks["expected_emergency_rows"]
assert abs(checks["exported_emergency_kwh"] - checks["model_emergency_kwh"]) < 1e-3
assert not errors
(ROOT / "04_q4/outputs/final/result4_2_export_validation.json").write_text(
    json.dumps(checks, ensure_ascii=False, indent=2), encoding="utf-8"
)
print(json.dumps(checks, ensure_ascii=False, indent=2))
